"""Petite bibliothèque de construction PPTX (python-pptx + XML) pour la soutenance Cléopâtre.
Chaque objet créé est enregistré (nom, géométrie, texte, couleur…) afin de vérifier
automatiquement la chaîne Morph (noms `!!` communs et réellement modifiés)."""
import copy, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from lxml import etree
from PIL import ImageFont, Image

W, H = 13.333, 7.5
INK, MUTED, LINE, LINE2 = "1D1D1F", "6E6E73", "D2D2D7", "AAAAAA"
GREEN, DEEP, CANVAS, WHITE = "075E46", "054537", "F5F5F7", "FFFFFF"
AMBER, GRAPHITE = "B25000", "2C2C2E"
ANTON, SANS, MONO = "Anton", "Instrument Sans", "JetBrains Mono"

FONTDIR = os.path.join(os.path.dirname(__file__), "polices")
_FILES = {(ANTON, False): "Anton-Regular.ttf", (ANTON, True): "Anton-Regular.ttf",
          (SANS, False): "InstrumentSans-Regular.ttf", (SANS, True): "InstrumentSans-Bold.ttf",
          (MONO, False): "JetBrainsMono-Regular.ttf", (MONO, True): "JetBrainsMono-Regular.ttf"}
_fc = {}
def _font(f, b):
    k = (f, b)
    if k not in _fc:
        _fc[k] = ImageFont.truetype(os.path.join(FONTDIR, _FILES[k]), 200)
    return _fc[k]
_cmap = {}
def _has(f, ch):
    from fontTools.ttLib import TTFont
    k = _FILES[(f, False)]
    if k not in _cmap:
        _cmap[k] = TTFont(os.path.join(FONTDIR, k)).getBestCmap()
    return ord(ch) in _cmap[k]

def rgb(h): return RGBColor.from_string(h)
def E(v): return int(round(v * 914400))

WARN = []

class Slide:
    def __init__(self, deck, idx, bg):
        self.deck, self.idx = deck, idx
        self.s = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])
        self.reg = {}
        self.sp = self.s.shapes
        self.s.background.fill.solid(); self.s.background.fill.fore_color.rgb = rgb(bg)
        self.bg = bg

    # ------------------------------------------------------------ helpers
    def _rec(self, name, **kw):
        if name in self.reg:
            raise ValueError(f"slide {self.idx}: nom en double {name}")
        self.reg[name] = kw

    @staticmethod
    def _strip_style(shape):
        st = shape._element.find(qn("p:style"))
        if st is not None: shape._element.remove(st)

    def _fill(self, shape, fill):
        if fill is None: shape.fill.background()
        else:
            shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)

    def _ln(self, shape, color, lw=0.75, dash=None, head=None, tail=None):
        ln = shape.line
        if color is None:
            ln.fill.background(); return
        ln.color.rgb = rgb(color); ln.width = Pt(lw)
        if dash: ln.dash_style = {"dash": MSO_LINE.DASH, "sysdash": MSO_LINE.SQUARE_DOT, "dot": MSO_LINE.ROUND_DOT}[dash]
        lnel = shape._element.spPr.find(qn("a:ln"))
        for tag, kind in (("a:headEnd", head), ("a:tailEnd", tail)):
            if kind:
                lnel.append(parse_xml(f'<{tag} xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" type="{kind}" w="med" len="med"/>'))

    # ------------------------------------------------------------ shapes
    def rect(self, name, x, y, w, h, fill=None, line=None, lw=0.75, dash=None, text=None, inset=0.12, **ts):
        shp = self.sp.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(h))
        self._strip_style(shp); shp.name = name
        self._fill(shp, fill); self._ln(shp, line, lw, dash)
        if text is not None:
            self._fill_text(shp, name, x, y, w, h, text, ts, inset=inset)
            txt = _plain(text)
        else:
            txt = ""
        self._rec(name, kind="rect", x=x, y=y, w=w, h=h, fill=fill, line=line, text=txt, size=ts.get("size"), shape=shp)
        return shp

    def line(self, name, x1, y1, x2, y2, color=INK, lw=1.0, dash=None, head=None, tail=None):
        c = self.sp.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
        self._strip_style(c); c.name = name
        self._ln(c, color, lw, dash, head, tail)
        self._rec(name, kind="line", x=min(x1, x2), y=min(y1, y2), w=abs(x2 - x1), h=abs(y2 - y1),
                  fill=None, line=color, text="", lw=lw, d=(x1 < x2, y1 < y2), dash=dash, shape=c)
        return c

    def tx(self, name, x, y, w, h, paras, **ts):
        tb = self.sp.add_textbox(E(x), E(y), E(w), E(h)); tb.name = name
        self._fill_text(tb, name, x, y, w, h, paras, ts, inset=0)
        self._rec(name, kind="text", x=x, y=y, w=w, h=h, fill=None, line=None, text=_plain(paras),
                  size=ts.get("size", 16), color=ts.get("color", INK), font=ts.get("font", SANS), shape=tb)
        return tb

    def pic(self, name, path, x, y, w, h, fx=0.5, fy=0.5, zoom=1.0):
        im = Image.open(path); iw, ih = im.size
        ir, br = iw / ih, w / h
        if br >= ir: vw, vh = 1.0, ir / br
        else: vh, vw = 1.0, br / ir
        vw, vh = vw / zoom, vh / zoom
        l = min(max(fx - vw / 2, 0), 1 - vw); t = min(max(fy - vh / 2, 0), 1 - vh)
        p = self.sp.add_picture(path, E(x), E(y), E(w), E(h)); p.name = name
        p.crop_left, p.crop_right = l, 1 - l - vw
        p.crop_top, p.crop_bottom = t, 1 - t - vh
        # texte alternatif
        p._element.nvPicPr.cNvPr.set("descr", os.path.basename(path))
        self._rec(name, kind="pic", x=x, y=y, w=w, h=h, fill=None, line=None, text="", crop=(round(l, 3), round(t, 3), round(vw, 3), round(vh, 3)), shape=p)
        return p

    # ------------------------------------------------------------ text engine
    def _fill_text(self, shp, name, x, y, w, h, paras, ts, inset=0.0):
        tf = shp.text_frame
        tf.word_wrap = True
        bp = tf._txBody.find(qn("a:bodyPr"))
        for k in ("lIns", "tIns", "rIns", "bIns"): bp.set(k, str(E(inset)))
        for ch in list(bp): bp.remove(ch)
        bp.append(parse_xml('<a:noAutofit xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'))
        anchor = ts.get("anchor", "t")
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        if isinstance(paras, (str, tuple)) or (isinstance(paras, list) and paras and isinstance(paras[0], tuple)):
            paras = [paras]
        dfont, dsize, dcol = ts.get("font", SANS), ts.get("size", 16), ts.get("color", INK)
        dbold, dalign, dls, dspc, dafter = ts.get("bold", False), ts.get("align", "l"), ts.get("ls"), ts.get("spc"), ts.get("after", 0)
        total_h = 0.0; measured = []
        first = True
        for p in paras:
            pd = {"runs": p} if not isinstance(p, dict) else dict(p)
            runs = pd["runs"]
            if isinstance(runs, (str, tuple)): runs = [runs]
            para = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
            para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[pd.get("align", dalign)]
            ls = pd.get("ls", dls)
            if ls: para.line_spacing = ls
            para.space_after = Pt(pd.get("after", dafter))
            toks = []
            psize = pd.get("size", dsize)
            for r in runs:
                if isinstance(r, str): r = (r, {})
                t, o = r
                f = o.get("font", pd.get("font", dfont)); sz = o.get("size", psize)
                b = o.get("bold", pd.get("bold", dbold)); c = o.get("color", pd.get("color", dcol))
                sp = o.get("spc", pd.get("spc", dspc))
                run = para.add_run(); run.text = t
                run.font.name = f; run.font.size = Pt(sz); run.font.bold = b; run.font.color.rgb = rgb(c)
                rPr = run._r.get_or_add_rPr(); rPr.set("lang", "fr-FR")
                if sp is not None: rPr.set("spc", str(int(sp * 100)))
                for ch in t:
                    if ch not in "\n" and not _has(f, ch): WARN.append(f"S{self.idx} {name}: glyphe absent {ch!r} ({f})")
                toks.append((t, f, sz, b, sp or 0))
            measured.append((toks, ls or 1.0, pd.get("after", dafter), psize, pd.get("font", dfont)))
        # mesure (approximation PIL)
        avail = (w - 2 * inset) * 72.0
        for toks, ls, after, psize, pfont in measured:
            lines, cur = 1, 0.0; maxw = 0.0; lh_pt = 0
            for t, f, sz, b, sp in toks:
                fo = _font(f, b)
                asc, desc = fo.getmetrics()
                lh_pt = max(lh_pt, (asc + desc) / 200.0 * sz)
                words = t.split(" ")
                for i, wd in enumerate(words):
                    ww = fo.getlength(wd) / 200.0 * sz + sp / 1.0 * len(wd) / 1.0 * 1.0
                    spw = fo.getlength(" ") / 200.0 * sz
                    add = ww + (spw if cur > 0 else 0)
                    if cur + add > avail + 0.5 and cur > 0:
                        lines += 1; cur = ww
                    else: cur += add
                    maxw = max(maxw, ww)
                    if ww > avail + 0.5: WARN.append(f"S{self.idx} {name}: mot trop large {wd!r} ({ww/72:.2f} in > {avail/72:.2f})")
            total_h += lines * lh_pt * ls + after
        if total_h / 72.0 > (h - 2 * inset) + 0.04 and anchor == "t" or total_h / 72.0 > h + 0.04:
            WARN.append(f"S{self.idx} {name}: texte trop haut ~{total_h/72:.2f} in > boîte {h:.2f} in")
        self.reg.setdefault("_fit", {})[name] = round(total_h / 72.0, 2)

    def notes(self, text):
        self.s.notes_slide.notes_text_frame.text = text

def _norm(p):
    if isinstance(p, (str, tuple)) or (isinstance(p, list) and p and isinstance(p[0], tuple)):
        return [p]
    return p
def _plain(p):
    out = []
    for q in _norm(p):
        runs = q["runs"] if isinstance(q, dict) else q
        if isinstance(runs, (str, tuple)): runs = [runs]
        out.append("".join(r if isinstance(r, str) else r[0] for r in runs))
    return "\n".join(out)

class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = E(W), E(H)
        self.slides = []
    def slide(self, bg=CANVAS):
        s = Slide(self, len(self.slides) + 1, bg); self.slides.append(s); return s

MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
def add_morph(slide, option="byObject", dur_ms=800):
    sld = slide.s._element
    xml = (f'<mc:AlternateContent xmlns:mc="{MC}" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
           f'<mc:Choice xmlns:p159="{P159}" xmlns:p14="{P14}" Requires="p159">'
           f'<p:transition spd="slow" p14:dur="{dur_ms}"><p159:morph option="{option}"/></p:transition></mc:Choice>'
           f'<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')
    el = parse_xml(xml)
    anchor = sld.find(qn("p:clrMapOvr"))
    if anchor is None: anchor = sld.find(qn("p:cSld"))
    anchor.addnext(el)
