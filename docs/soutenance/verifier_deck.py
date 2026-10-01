# -*- coding: utf-8 -*-
"""Vérifie le .pptx exporté : 20 diapositives, Morph natif sur 2→20, noms !! communs et réellement modifiés,
cadres de diagrammes vides, images non déformées."""
import sys, re, zipfile
from pptx import Presentation
from pptx.util import Emu
from lxml import etree
path = sys.argv[1] if len(sys.argv) > 1 else "Cleopatre-PFE-Soutenance.pptx"
p = Presentation(path)
n = len(p.slides); assert n == 20, n
print("diapositives :", n)
ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main", "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p159": "http://schemas.microsoft.com/office/powerpoint/2015/09/main", "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006"}
def props(sh):
    el = sh._element
    xfrm = el.find(".//a:xfrm", ns)
    if xfrm is None: xfrm = el.find(".//p:xfrm", ns)
    g = tuple(xfrm.find(t, ns).attrib.get(k) for t, k in (("a:off", "x"), ("a:off", "y"), ("a:ext", "cx"), ("a:ext", "cy"))) + (xfrm.get("rot"), xfrm.get("flipH"), xfrm.get("flipV"))
    txt = "".join(el.itertext()).strip() if sh.has_text_frame else ""
    fill = tuple(etree.tostring(f) for f in el.findall(".//p:spPr/a:solidFill", ns))
    sz = tuple(r.get("sz") for r in el.findall(".//a:rPr", ns))
    crop = etree.tostring(el.find(".//a:srcRect", ns)) if el.find(".//a:srcRect", ns) is not None else b""
    return dict(geom=g, text=txt, fill=fill, size=sz, crop=crop)
names = []
ok = True
for i, s in enumerate(p.slides, 1):
    d = {}
    for sh in s.shapes:
        if sh.name.startswith("!!"):
            assert sh.name not in d, f"S{i} doublon {sh.name}"
            d[sh.name] = props(sh)
    names.append(d)
    tr = s._element.find("mc:AlternateContent", ns)
    has = tr is not None and tr.find(".//p159:morph", ns) is not None
    opt = tr.find(".//p159:morph", ns).get("option") if has else None
    dur = tr.find(".//{http://schemas.microsoft.com/office/powerpoint/2010/main}dummy", ns) if False else None
    if i == 1: assert not has, "S1 ne doit pas avoir de transition"
    else:
        if not has: ok = False; print(f"✗ S{i}: pas de Morph")
for i in range(1, n):
    a, b = names[i - 1], names[i]
    common = sorted(set(a) & set(b))
    changed = [k for k in common if a[k] != b[k]]
    status = "✓" if len(common) >= 2 and len(changed) >= 2 else "✗"
    if status == "✗": ok = False
    print(f"{status} {i:>2}→{i+1:<2} communs={len(common):>2} modifiés={len(changed):>2}  clés: {', '.join(k[2:] for k in changed if not re.match(r'(Node|Link|Corner)', k[2:]))}")
# diagrammes vides
for i in range(11, 16):
    s = p.slides[i - 1]
    inside = [sh.name for sh in s.shapes if sh.name == "!!DiagramFrame" and sh.has_text_frame and sh.text_frame.text.strip()]
    assert not inside
    pics = [sh.name for sh in s.shapes if sh.shape_type == 13]
    assert not pics, pics
print("cadres de diagramme vides : ✓ (aucune image/texte dans !!DiagramFrame, S11–15)")
# images non déformées
from PIL import Image
import io
for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        if sh.shape_type == 13:
            im = Image.open(io.BytesIO(sh.image.blob)); iw, ih = im.size
            vw = iw * (1 - sh.crop_left - sh.crop_right); vh = ih * (1 - sh.crop_top - sh.crop_bottom)
            r1, r2 = vw / vh, sh.width / sh.height
            flag = "✓" if abs(r1 / r2 - 1) < 0.01 else "✗"
            if flag == "✗": ok = False
            print(f"{flag} S{i} {sh.name:<16} ratio image {r1:.3f} / cadre {r2:.3f}")
z = zipfile.ZipFile(path)
ntr = sum(1 for f in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml", f) and b"p159:morph" in z.read(f))
print("transitions Morph natives :", ntr)
print("RÉSULTAT :", "OK" if ok and ntr == 19 else "ÉCHEC")
